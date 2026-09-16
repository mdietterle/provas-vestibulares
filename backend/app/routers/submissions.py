import difflib
import io
import math
import os
import statistics
import uuid
from datetime import datetime
from typing import List

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload, selectinload

from app.core.config import settings
from app.database import SessionLocal, get_db
from app.deps import get_current_user, require_professor
from app.models import (
    Exam,
    ExamQuestion,
    ExamSubmission,
    Question,
    QuestionOption,
    QuestionType,
    StudentClass,
    Subject,
    SubmissionAnswer,
    SubmissionStatus,
    User,
    UserRole,
)
from app.schemas import (
    AnswerScoreOverride,
    SubmissionCreate,
    SubmissionListOut,
    SubmissionOut,
)
from app.services.ai_corrector import correct_essay, correct_essay_from_image, groq_is_available, ollama_is_available
from app.services import usage as usage_service
from app.models import UsageEventType

router = APIRouter(prefix="/submissions", tags=["submissions"])


def _assert_exam_manage_access(db: Session, exam: Exam, current_user: User):
    """Only professors/admins from the exam's own institution may manage its submissions."""
    if not exam:
        raise HTTPException(404, "Prova não encontrada")
    subject = db.get(Subject, exam.subject_id)
    if not subject or subject.institution_id != current_user.institution_id:
        raise HTTPException(403, "Sem acesso")
    if current_user.role == UserRole.PROFESSOR and exam.professor_id != current_user.id:
        raise HTTPException(403, "Sem acesso")


def _assert_student_can_submit(db: Session, exam: Exam, current_user: User):
    """Only students enrolled in the exam's own class/institution may submit to it."""
    if not exam:
        raise HTTPException(404, "Prova não encontrada")
    subject = db.get(Subject, exam.subject_id)
    if not subject or subject.institution_id != current_user.institution_id:
        raise HTTPException(403, "Sem acesso")
    enrolled = (
        db.query(StudentClass)
        .filter(StudentClass.student_id == current_user.id, StudentClass.class_id == exam.class_id)
        .first()
    )
    if not enrolled:
        raise HTTPException(403, "Você não está matriculado na turma desta prova")


def _ai_correction_allowed(db: Session, professor_id: int) -> bool:
    """True se o plano da instituição do professor inclui correção por IA."""
    prof = db.get(User, professor_id)
    if not prof:
        return False
    return usage_service.ai_enabled(db, prof.institution_id)


def _check_ai_correction_quota(db: Session, professor_id: int) -> dict:
    """Retorna resultado de check_quota para correção do professor."""
    prof = db.get(User, professor_id)
    if not prof:
        return {"allowed": False, "used": 0, "limit": 0, "remaining": 0, "overage_count": 0}
    return usage_service.check_quota(db, prof, UsageEventType.AI_CORRECTION.value, quantity=1)


def _record_ai_correction(db: Session, professor_id: int):
    """Contabiliza uma discursiva corrigida por IA, atribuída ao professor."""
    prof = db.get(User, professor_id)
    if prof:
        usage_service.record_usage(db, prof, UsageEventType.AI_CORRECTION.value, quantity=1)

SCANS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "uploads", "scans")
try:
    os.makedirs(SCANS_DIR, exist_ok=True)
except Exception:
    pass


def _load_submission_full(db: Session, submission_id: int):
    """Load submission with all nested relationships in 3 queries instead of N*3."""
    return (
        db.query(ExamSubmission)
        .filter(ExamSubmission.id == submission_id)
        .options(
            selectinload(ExamSubmission.answers)
            .joinedload(SubmissionAnswer.exam_question)
            .joinedload(ExamQuestion.question)
        )
        .first()
    )


def _do_correct(submission_id: int):
    """Run correction in a background task with its own DB session."""
    db = SessionLocal()
    try:
        submission = _load_submission_full(db, submission_id)
        if not submission:
            return

        submission.status = SubmissionStatus.CORRECTING
        submission.correcting_since = datetime.utcnow()
        db.commit()

        exam = db.get(Exam, submission.exam_id)
        professor_id = exam.professor_id if exam else None
        ai_ok = professor_id is not None and _ai_correction_allowed(db, professor_id)

        total = 0.0
        all_scored = True

        for answer in submission.answers:
            eq: ExamQuestion = answer.exam_question
            q: Question = eq.question

            if q.question_type in (QuestionType.MULTIPLE_CHOICE, QuestionType.TRUE_FALSE):
                if answer.selected_option_id:
                    opt = db.get(QuestionOption, answer.selected_option_id)
                    answer.score = eq.points if (opt and opt.is_correct) else 0.0
                else:
                    answer.score = 0.0
                answer.is_auto_corrected = True

            elif q.question_type == QuestionType.SUMMATION:
                correct_sum = sum(o.order for o in q.options if o.is_correct)
                try:
                    student_sum = int(answer.essay_text or "0")
                except (ValueError, TypeError):
                    student_sum = -1
                answer.score = eq.points if student_sum == correct_sum else 0.0
                answer.is_auto_corrected = True

            elif q.question_type == QuestionType.ESSAY:
                if not ai_ok:
                    answer.ai_feedback = (
                        "Correção por IA indisponível no plano atual. Revisão manual necessária."
                    )
                    answer.score = None
                    all_scored = False
                    continue
                prof = db.get(User, professor_id)
                quota = usage_service.check_quota(
                    db, prof, UsageEventType.AI_CORRECTION.value, quantity=1
                ) if prof else {"allowed": False}
                if not quota["allowed"]:
                    answer.ai_feedback = (
                        "Cota mensal de correções por IA esgotada. Revisão manual necessária."
                    )
                    answer.score = None
                    all_scored = False
                    continue
                try:
                    if answer.essay_image_base64:
                        result = correct_essay_from_image(
                            question_statement=q.statement,
                            criteria=q.criteria,
                            image_base64=answer.essay_image_base64,
                            max_points=eq.points,
                        )
                        # Salva transcrição como essay_text para o professor ver
                        if result.get("transcription"):
                            answer.essay_text = result["transcription"]
                    else:
                        result = correct_essay(
                            question_statement=q.statement,
                            criteria=q.criteria,
                            student_answer=answer.essay_text or "",
                            max_points=eq.points,
                        )
                    answer.score = result["score"]
                    answer.ai_feedback = result["feedback"]
                    answer.is_auto_corrected = True
                    _record_ai_correction(db, professor_id)
                except Exception as e:
                    answer.ai_feedback = f"Erro na IA: {e}. Revisão manual necessária."
                    answer.score = None
                    all_scored = False

            if answer.score is not None:
                total += answer.score

        submission.status = SubmissionStatus.DONE if all_scored else SubmissionStatus.CORRECTING
        submission.total_score = round(total, 2) if all_scored else None
        if all_scored:
            submission.correcting_since = None
        db.commit()
    finally:
        db.close()


# ── Submit (student only — no auto-correction) ────────────────────────────────

@router.post("/exams/{exam_id}", response_model=SubmissionOut, status_code=201)
def submit_exam(
    exam_id: int,
    payload: SubmissionCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.role != UserRole.STUDENT:
        raise HTTPException(403, "Apenas alunos podem submeter provas")

    exam = db.get(Exam, exam_id)
    _assert_student_can_submit(db, exam, current_user)

    if db.query(ExamSubmission).filter_by(exam_id=exam_id, student_id=current_user.id).first():
        raise HTTPException(400, "Você já submeteu esta prova")

    submission = ExamSubmission(
        exam_id=exam_id,
        student_id=current_user.id,
        status=SubmissionStatus.PENDING,
    )
    db.add(submission)
    db.flush()

    eq_map = {eq.id: eq for eq in exam.exam_questions}
    for ans in payload.answers:
        if ans.exam_question_id not in eq_map:
            raise HTTPException(400, f"Questão {ans.exam_question_id} não pertence a esta prova")
        db.add(SubmissionAnswer(
            submission_id=submission.id,
            exam_question_id=ans.exam_question_id,
            selected_option_id=ans.selected_option_id,
            essay_text=ans.essay_text,
            essay_image_base64=ans.essay_image_base64,
        ))

    db.commit()
    db.refresh(submission)
    return submission


# ── Trigger correction (professor) ────────────────────────────────────────────

@router.post("/{submission_id}/correct", response_model=SubmissionOut)
def correct_submission(
    submission_id: int,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    sub = _load_submission_full(db, submission_id)
    if not sub:
        raise HTTPException(404, "Submissão não encontrada")

    exam = db.get(Exam, sub.exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    has_essay = any(
        a.exam_question.question.question_type == QuestionType.ESSAY
        for a in sub.answers
    )
    if has_essay and not ollama_is_available():
        raise HTTPException(503, "GROQ_API_KEY não configurada. Defina a variável no .env para habilitar correção por IA.")

    sub.status = SubmissionStatus.CORRECTING
    sub.correcting_since = datetime.utcnow()
    db.commit()
    db.refresh(sub)
    background_tasks.add_task(_do_correct, sub.id)
    return sub


# ── Correct all submitted submissions for an exam (professor) ────────────────

def _do_correct_all(submission_ids: list):
    """Correct submissions one by one to avoid overloading the Groq API."""
    for sid in submission_ids:
        try:
            _do_correct(sid)
        except Exception as e:
            print(f"Bulk correction error for submission {sid}: {e}")


@router.post("/exams/{exam_id}/correct-all")
def correct_all_submissions(
    exam_id: int,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    exam = db.get(Exam, exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    subs = (
        db.query(ExamSubmission)
        .filter_by(exam_id=exam_id, status=SubmissionStatus.PENDING)
        .all()
    )
    if not subs:
        raise HTTPException(400, "Nenhuma submissão pendente de correção")

    has_essay = (
        db.query(ExamQuestion)
        .join(ExamQuestion.question)
        .filter(
            ExamQuestion.exam_id == exam_id,
            Question.question_type == QuestionType.ESSAY,
        )
        .first()
        is not None
    )
    if has_essay and not ollama_is_available():
        raise HTTPException(
            503,
            "GROQ_API_KEY não configurada. Defina a variável no .env para habilitar correção por IA.",
        )

    now = datetime.utcnow()
    for sub in subs:
        sub.status = SubmissionStatus.CORRECTING
        sub.correcting_since = now
    db.commit()

    submission_ids = [sub.id for sub in subs]
    background_tasks.add_task(_do_correct_all, submission_ids)

    return {"queued": len(submission_ids), "submission_ids": submission_ids}


# ── Manually finalize a stuck submission (professor) ─────────────────────────

@router.post("/{submission_id}/finalize", response_model=SubmissionOut)
def finalize_submission(
    submission_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    """Finalize correction: auto-scores any pending MC/TF answers, blocks only on unscored essays."""
    sub = _load_submission_full(db, submission_id)
    if not sub:
        raise HTTPException(404, "Submissão não encontrada")
    exam = db.get(Exam, sub.exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    unscored_essays = []
    for answer in sub.answers:
        if answer.score is not None:
            continue
        q = answer.exam_question.question
        if q.question_type in (QuestionType.MULTIPLE_CHOICE, QuestionType.TRUE_FALSE):
            if answer.selected_option_id:
                opt = db.get(QuestionOption, answer.selected_option_id)
                answer.score = answer.exam_question.points if (opt and opt.is_correct) else 0.0
            else:
                answer.score = 0.0
            answer.is_auto_corrected = True
        elif q.question_type == QuestionType.SUMMATION:
            correct_sum = sum(o.order for o in q.options if o.is_correct)
            try:
                student_sum = int(answer.essay_text or "0")
            except (ValueError, TypeError):
                student_sum = -1
            answer.score = answer.exam_question.points if student_sum == correct_sum else 0.0
            answer.is_auto_corrected = True
        else:
            unscored_essays.append(answer)

    if unscored_essays:
        raise HTTPException(
            400,
            f"{len(unscored_essays)} questão(ões) dissertativa(s) ainda sem nota. "
            "Atribua as notas manualmente antes de finalizar."
        )

    sub.total_score = round(sum(a.score for a in sub.answers), 2)
    sub.status = SubmissionStatus.DONE
    sub.correcting_since = None
    db.commit()
    db.refresh(sub)
    return sub


# ── Release a single submission's grade (professor) ───────────────────────────

@router.post("/{submission_id}/release", response_model=SubmissionOut)
def release_submission(
    submission_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    sub = _load_submission_full(db, submission_id)
    if not sub:
        raise HTTPException(404, "Submissão não encontrada")
    exam = db.get(Exam, sub.exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    if sub.status != SubmissionStatus.DONE:
        raise HTTPException(400, "Só é possível liberar submissões já corrigidas")

    sub.status = SubmissionStatus.RELEASED
    db.commit()
    db.refresh(sub)
    return sub


# ── Release all corrected grades for an exam (professor) ─────────────────────

@router.post("/exams/{exam_id}/release-all", response_model=List[SubmissionListOut])
def release_all_submissions(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    exam = db.get(Exam, exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    subs = (
        db.query(ExamSubmission)
        .filter_by(exam_id=exam_id)
        .filter(ExamSubmission.status == SubmissionStatus.DONE)
        .all()
    )
    if not subs:
        raise HTTPException(400, "Nenhuma prova corrigida aguardando liberação")

    for sub in subs:
        sub.status = SubmissionStatus.RELEASED
    db.commit()

    return (
        db.query(ExamSubmission)
        .filter_by(exam_id=exam_id)
        .options(joinedload(ExamSubmission.student))
        .all()
    )


# ── List submissions by exam (professor) ──────────────────────────────────────

@router.get("/exams/{exam_id}", response_model=List[SubmissionListOut])
def list_submissions(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    exam = db.get(Exam, exam_id)
    _assert_exam_manage_access(db, exam, current_user)
    return (
        db.query(ExamSubmission)
        .filter_by(exam_id=exam_id)
        .options(joinedload(ExamSubmission.student))
        .all()
    )


# ── Export grades as XLSX (professor) ────────────────────────────────────────

@router.get("/exams/{exam_id}/export")
def export_grades_xlsx(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    exam = (
        db.query(Exam)
        .options(selectinload(Exam.exam_questions), joinedload(Exam.class_))
        .filter(Exam.id == exam_id)
        .first()
    )
    _assert_exam_manage_access(db, exam, current_user)

    students = (
        db.query(User)
        .join(StudentClass, StudentClass.student_id == User.id)
        .filter(StudentClass.class_id == exam.class_id)
        .order_by(User.name)
        .all()
    )

    sub_map = {
        s.student_id: s
        for s in db.query(ExamSubmission).filter_by(exam_id=exam_id).all()
    }

    total_pts = sum(eq.points for eq in exam.exam_questions)

    STATUS_LABELS = {
        "pending": "Pendente",
        "correcting": "Corrigindo",
        "done": "Concluído",
        "released": "Liberado",
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"

    header_fill = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center")

    headers = ["Aluno", "Turma", "Ano", "Status", "Nota", "Nota Máxima"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    class_name = exam.class_.name if exam.class_ else ""
    class_year = exam.class_.year if exam.class_ else ""

    for student in students:
        sub = sub_map.get(student.id)
        nota = sub.total_score if sub and sub.total_score is not None else ""
        status = STATUS_LABELS.get(sub.status, "—") if sub else "Não entregou"
        ws.append([student.name, class_name, class_year, status, nota, total_pts])

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_title = exam.title.replace(" ", "_")[:40]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="notas_{safe_title}.xlsx"'},
    )


# ── Get my submission (student) ───────────────────────────────────────────────

@router.get("/my/{exam_id}", response_model=SubmissionOut)
def my_submission(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    sub = (
        db.query(ExamSubmission)
        .filter_by(exam_id=exam_id, student_id=current_user.id)
        .first()
    )
    if not sub:
        raise HTTPException(404, "Você ainda não entregou esta prova")
    if sub.status != SubmissionStatus.RELEASED:
        raise HTTPException(403, "A nota ainda não foi disponibilizada pelo professor")
    return sub


# ── Get full submission detail (professor) ────────────────────────────────────

@router.get("/{submission_id}", response_model=SubmissionOut)
def get_submission(
    submission_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    sub = _load_submission_full(db, submission_id)
    if not sub:
        raise HTTPException(404, "Submissão não encontrada")
    exam = db.get(Exam, sub.exam_id)
    _assert_exam_manage_access(db, exam, current_user)
    return sub


# ── Override score for an answer (professor) ──────────────────────────────────

@router.patch("/{submission_id}/answers/{answer_id}", response_model=SubmissionOut)
def override_score(
    submission_id: int,
    answer_id: int,
    payload: AnswerScoreOverride,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    sub = _load_submission_full(db, submission_id)
    if not sub:
        raise HTTPException(404, "Submissão não encontrada")
    exam = db.get(Exam, sub.exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    answer = db.get(SubmissionAnswer, answer_id)
    if not answer or answer.submission_id != submission_id:
        raise HTTPException(404, "Resposta não encontrada")

    answer.score = payload.score
    if payload.ai_feedback is not None:
        answer.ai_feedback = payload.ai_feedback

    # sub.answers already loaded — no extra queries
    scores = [a.score for a in sub.answers]
    if all(s is not None for s in scores):
        sub.total_score = round(sum(scores), 2)  # type: ignore[arg-type]
        if sub.status == SubmissionStatus.CORRECTING:
            sub.status = SubmissionStatus.DONE
            sub.correcting_since = None

    db.commit()
    db.refresh(sub)
    return sub


# ── Re-run AI correction for a single answer (professor) ─────────────────────

@router.post("/{submission_id}/answers/{answer_id}/ai-correct", response_model=SubmissionOut)
def ai_correct_answer(
    submission_id: int,
    answer_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    sub = db.get(ExamSubmission, submission_id)
    if not sub:
        raise HTTPException(404, "Submissão não encontrada")
    exam = db.get(Exam, sub.exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    answer = db.get(SubmissionAnswer, answer_id)
    if not answer or answer.submission_id != submission_id:
        raise HTTPException(404, "Resposta não encontrada")

    eq: ExamQuestion = answer.exam_question
    q: Question = eq.question

    if q.question_type != QuestionType.ESSAY:
        raise HTTPException(400, "Correção por IA disponível apenas para questões dissertativas")

    if not ollama_is_available():
        raise HTTPException(503, "GROQ_API_KEY não configurada.")

    if not _ai_correction_allowed(db, exam.professor_id):
        raise HTTPException(
            403,
            "Correção de discursivas com IA não está disponível no plano Basic. "
            "Faça upgrade para o plano Pro ou Enterprise.",
        )

    quota = _check_ai_correction_quota(db, exam.professor_id)
    if not quota["allowed"]:
        raise HTTPException(
            402,
            {
                "code": "quota_exceeded",
                "resource": "ai_correction",
                "used": quota["used"],
                "limit": quota["limit"],
                "remaining": quota["remaining"],
                "message": (
                    f"Você utilizou {quota['used']} de {quota['limit']} correções por IA "
                    f"disponíveis neste mês. Adquira um pacote avulso para continuar."
                ),
            },
        )

    try:
        if answer.essay_image_base64:
            result = correct_essay_from_image(
                question_statement=q.statement,
                criteria=q.criteria,
                image_base64=answer.essay_image_base64,
                max_points=eq.points,
            )
            if result.get("transcription"):
                answer.essay_text = result["transcription"]
        else:
            result = correct_essay(
                question_statement=q.statement,
                criteria=q.criteria,
                student_answer=answer.essay_text or "",
                max_points=eq.points,
            )
        answer.score = result["score"]
        answer.ai_feedback = result["feedback"]
        answer.is_auto_corrected = True
        _record_ai_correction(db, exam.professor_id)
    except Exception as e:
        raise HTTPException(500, f"Erro na correção por IA: {e}")

    all_scored = all(a.score is not None for a in sub.answers)
    if all_scored:
        sub.total_score = round(sum(a.score for a in sub.answers), 2)
        sub.status = SubmissionStatus.DONE
        sub.correcting_since = None

    db.commit()
    db.refresh(sub)
    return sub


# ── Upload scanned exam (student) ─────────────────────────────────────────────

def _do_scan_correct(submission_id: int, image_path: str):
    """Background task: correct a scanned exam using Ollama vision."""
    from app.services.scan_corrector import correct_from_scan
    db = SessionLocal()
    try:
        submission = db.get(ExamSubmission, submission_id)
        if not submission:
            return
        with open(image_path, "rb") as f:
            image_bytes = f.read()
        correct_from_scan(submission, image_bytes, db)
    except Exception as e:
        print(f"Scan correction error for submission {submission_id}: {e}")
    finally:
        db.close()


@router.post("/exams/{exam_id}/upload-scan", response_model=SubmissionOut, status_code=201)
async def upload_scan(
    exam_id: int,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.role != UserRole.STUDENT:
        raise HTTPException(403, "Apenas alunos podem enviar provas digitalizadas")

    exam = db.get(Exam, exam_id)
    _assert_student_can_submit(db, exam, current_user)

    # Validate file type
    allowed = {"image/jpeg", "image/png", "image/webp", "application/pdf"}
    if file.content_type not in allowed:
        raise HTTPException(400, "Formato inválido. Envie JPEG, PNG, WebP ou PDF.")

    # Read file content
    content = await file.read()
    if len(content) > 20 * 1024 * 1024:  # 20 MB limit
        raise HTTPException(400, "Arquivo muito grande. Limite: 20 MB.")

    # Get or create submission
    sub = db.query(ExamSubmission).filter_by(exam_id=exam_id, student_id=current_user.id).first()
    if sub and sub.status == SubmissionStatus.RELEASED:
        raise HTTPException(400, "Nota já liberada. Não é possível reenviar.")

    # Save file
    ext = file.filename.rsplit(".", 1)[-1].lower() if file.filename and "." in file.filename else "jpg"
    filename = f"{uuid.uuid4().hex}.{ext}"
    file_path = os.path.join(SCANS_DIR, filename)
    with open(file_path, "wb") as f:
        f.write(content)

    if not sub:
        sub = ExamSubmission(
            exam_id=exam_id,
            student_id=current_user.id,
            status=SubmissionStatus.CORRECTING,
            correcting_since=datetime.utcnow(),
            scan_file=filename,
        )
        db.add(sub)
    else:
        # Delete old scan file if exists
        if sub.scan_file:
            old_path = os.path.join(SCANS_DIR, sub.scan_file)
            if os.path.exists(old_path):
                os.remove(old_path)
        sub.scan_file = filename
        sub.status = SubmissionStatus.CORRECTING
        sub.correcting_since = datetime.utcnow()

    db.commit()
    db.refresh(sub)

    background_tasks.add_task(_do_scan_correct, sub.id, file_path)
    return sub


# ── Professor uploads scanned exam (identifies student automatically) ─────────

@router.post("/exams/{exam_id}/professor-upload-scan", status_code=201)
async def professor_upload_scan(
    exam_id: int,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    student_id: int = Form(0),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    from app.models import StudentClass
    from app.services.scan_corrector import identify_student_from_scan

    exam = db.get(Exam, exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    allowed = {"image/jpeg", "image/png", "image/webp", "application/pdf"}
    if file.content_type not in allowed:
        raise HTTPException(400, "Formato inválido. Envie JPEG, PNG, WebP ou PDF.")

    content = await file.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(400, "Arquivo muito grande. Limite: 20 MB.")

    # Fetch enrolled students for identification / validation
    enrolled_ids = [
        r[0] for r in
        db.query(StudentClass.student_id).filter(StudentClass.class_id == exam.class_id).all()
    ]
    enrolled_students = db.query(User).filter(User.id.in_(enrolled_ids)).all()

    resolved_student_id = student_id if student_id else None
    identification: dict = {}

    if not resolved_student_id:
        from app.services.scan_corrector import ollama_vision_available
        if not ollama_vision_available():
            raise HTTPException(503, "GROQ_API_KEY não configurada. Selecione o aluno manualmente.")
        identification = identify_student_from_scan(content, exam_id, enrolled_students)
        resolved_student_id = identification.get("student_id")

        if not resolved_student_id:
            # Return 422 so frontend can ask professor to select manually
            raise HTTPException(
                422,
                {
                    "message": "Não foi possível identificar o aluno automaticamente.",
                    "read_name": identification.get("read_name") or identification.get("student_name"),
                    "method": identification.get("method"),
                },
            )

    # Validate student belongs to the exam's class
    if resolved_student_id not in enrolled_ids:
        raise HTTPException(400, "Aluno não matriculado nesta turma.")

    # Save file
    ext = (file.filename or "scan.jpg").rsplit(".", 1)[-1].lower()
    filename = f"{uuid.uuid4().hex}.{ext}"
    file_path = os.path.join(SCANS_DIR, filename)
    with open(file_path, "wb") as f:
        f.write(content)

    # Create or replace submission
    sub = db.query(ExamSubmission).filter_by(exam_id=exam_id, student_id=resolved_student_id).first()
    if sub and sub.status == SubmissionStatus.RELEASED:
        raise HTTPException(400, "Nota já liberada para este aluno. Não é possível substituir.")

    if sub:
        if sub.scan_file:
            old_path = os.path.join(SCANS_DIR, sub.scan_file)
            if os.path.exists(old_path):
                os.remove(old_path)
        sub.scan_file = filename
        sub.status = SubmissionStatus.CORRECTING
        sub.correcting_since = datetime.utcnow()
    else:
        sub = ExamSubmission(
            exam_id=exam_id,
            student_id=resolved_student_id,
            status=SubmissionStatus.CORRECTING,
            correcting_since=datetime.utcnow(),
            scan_file=filename,
        )
        db.add(sub)

    db.commit()
    db.refresh(sub)
    background_tasks.add_task(_do_scan_correct, sub.id, file_path)

    # Return submission + identification metadata
    from app.schemas import SubmissionOut
    result = SubmissionOut.model_validate(sub)
    return {
        **result.model_dump(),
        "identified_student": identification.get("student_name") or sub.student.name,
        "identification_method": identification.get("method", "manual"),
        "identification_confidence": identification.get("confidence", "high"),
    }


@router.get("/exams/{exam_id}/analytics")
def exam_analytics(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_professor),
):
    exam = db.get(Exam, exam_id)
    _assert_exam_manage_access(db, exam, current_user)

    # Load all done/released submissions with answers
    subs = (
        db.query(ExamSubmission)
        .filter(ExamSubmission.exam_id == exam_id)
        .filter(ExamSubmission.status.in_([SubmissionStatus.DONE, SubmissionStatus.RELEASED]))
        .options(
            selectinload(ExamSubmission.answers)
            .joinedload(SubmissionAnswer.exam_question)
            .joinedload(ExamQuestion.question)
            .selectinload(Question.options),
            joinedload(ExamSubmission.student),
        )
        .all()
    )

    total_enrolled = (
        db.query(StudentClass).filter(StudentClass.class_id == exam.class_id).count()
    )
    total_submitted = (
        db.query(ExamSubmission).filter(ExamSubmission.exam_id == exam_id).count()
    )
    total_points = sum(eq.points for eq in exam.exam_questions)

    # ── Grade stats ───────────────────────────────────────────────────────────
    scores = [s.total_score for s in subs if s.total_score is not None]
    avg_score = round(statistics.mean(scores), 2) if scores else None
    median_score = round(statistics.median(scores), 2) if scores else None
    highest = round(max(scores), 2) if scores else None
    lowest = round(min(scores), 2) if scores else None
    std_dev = round(statistics.stdev(scores), 2) if len(scores) >= 2 else 0.0
    pass_threshold = total_points * 0.5
    pass_rate = round(sum(1 for s in scores if s >= pass_threshold) / len(scores) * 100, 1) if scores else 0.0

    # Score distribution in 5 buckets (0-20%, 20-40%, ...)
    buckets = [0, 0, 0, 0, 0]
    bucket_labels = ["0–20%", "20–40%", "40–60%", "60–80%", "80–100%"]
    for score in scores:
        pct = (score / total_points) if total_points > 0 else 0
        idx = min(int(pct * 5), 4)
        buckets[idx] += 1
    score_distribution = [{"label": label, "count": count} for label, count in zip(bucket_labels, buckets)]

    # ── Per-question stats ────────────────────────────────────────────────────
    eq_map: dict[int, dict] = {}
    for eq in exam.exam_questions:
        eq_map[eq.id] = {
            "exam_question_id": eq.id,
            "order": eq.order + 1,
            "statement": eq.question.statement,
            "question_type": eq.question.question_type,
            "points": eq.points,
            "correct_count": 0,
            "wrong_count": 0,
            "skip_count": 0,
            "score_sum": 0.0,
            "answer_count": 0,
        }

    for sub in subs:
        for ans in sub.answers:
            stat = eq_map.get(ans.exam_question_id)
            if not stat:
                continue
            stat["answer_count"] += 1
            q = ans.exam_question.question
            if q.question_type in (QuestionType.MULTIPLE_CHOICE, QuestionType.TRUE_FALSE):
                if ans.selected_option_id is None:
                    stat["skip_count"] += 1
                else:
                    opt = next((o for o in q.options if o.id == ans.selected_option_id), None)
                    if opt and opt.is_correct:
                        stat["correct_count"] += 1
                    else:
                        stat["wrong_count"] += 1
            else:
                if ans.score is not None:
                    stat["score_sum"] += ans.score
                    if ans.score >= ans.exam_question.points:
                        stat["correct_count"] += 1
                    elif ans.score == 0:
                        stat["wrong_count"] += 1
                    else:
                        stat["wrong_count"] += 1  # partial = wrong for error rate

    question_stats = []
    for stat in sorted(eq_map.values(), key=lambda x: x["order"]):
        total_answers = stat["correct_count"] + stat["wrong_count"] + stat["skip_count"]
        error_rate = round(stat["wrong_count"] / total_answers, 3) if total_answers > 0 else 0.0
        avg_q_score = round(stat["score_sum"] / stat["answer_count"], 2) if stat["answer_count"] > 0 else None
        question_stats.append({
            "exam_question_id": stat["exam_question_id"],
            "order": stat["order"],
            "statement": stat["statement"],
            "question_type": stat["question_type"],
            "points": stat["points"],
            "correct_count": stat["correct_count"],
            "wrong_count": stat["wrong_count"],
            "skip_count": stat["skip_count"],
            "avg_score": avg_q_score,
            "error_rate": error_rate,
        })

    # ── Plagiarism detection ──────────────────────────────────────────────────
    plagiarism_alerts = []

    # MC/TF: flag pairs with ≥3 matching wrong answers on same wrong option
    mc_questions = [eq for eq in exam.exam_questions
                    if eq.question.question_type in (QuestionType.MULTIPLE_CHOICE, QuestionType.TRUE_FALSE)]

    if mc_questions:
        # Build map: student_id → {exam_question_id: selected_option_id (if wrong)}
        wrong_answers: dict[int, dict[int, int]] = {}
        for sub in subs:
            wrong_by_q: dict[int, int] = {}
            for ans in sub.answers:
                q = ans.exam_question.question
                if q.question_type not in (QuestionType.MULTIPLE_CHOICE, QuestionType.TRUE_FALSE):
                    continue
                if ans.selected_option_id is None:
                    continue
                opt = next((o for o in q.options if o.id == ans.selected_option_id), None)
                if opt and not opt.is_correct:
                    wrong_by_q[ans.exam_question_id] = ans.selected_option_id
            wrong_answers[sub.student_id] = wrong_by_q

        student_list = [(sub.student_id, sub.student.name) for sub in subs]
        for i in range(len(student_list)):
            for j in range(i + 1, len(student_list)):
                id_a, name_a = student_list[i]
                id_b, name_b = student_list[j]
                wa, wb = wrong_answers.get(id_a, {}), wrong_answers.get(id_b, {})
                matching = [(qid, opt) for qid, opt in wa.items() if wb.get(qid) == opt]
                if len(matching) >= 3:
                    similarity = round(len(matching) / max(len(mc_questions), 1), 2)
                    q_labels = []
                    for qid, _ in matching[:5]:
                        eq = next((e for e in mc_questions if e.id == qid), None)
                        if eq:
                            q_labels.append(f"Q{eq.order + 1}")
                    plagiarism_alerts.append({
                        "student_a": {"id": id_a, "name": name_a},
                        "student_b": {"id": id_b, "name": name_b},
                        "type": "multiple_choice",
                        "similarity": similarity,
                        "evidence": f"{len(matching)} respostas erradas idênticas: {', '.join(q_labels)}",
                    })

    # Essay: flag pairs with text similarity > 0.65
    essay_questions = [eq for eq in exam.exam_questions
                       if eq.question.question_type == QuestionType.ESSAY]

    if essay_questions:
        essay_texts: dict[int, dict[int, str]] = {}
        for sub in subs:
            texts_by_q: dict[int, str] = {}
            for ans in sub.answers:
                if ans.exam_question.question.question_type == QuestionType.ESSAY and ans.essay_text:
                    texts_by_q[ans.exam_question_id] = ans.essay_text.strip()
            essay_texts[sub.student_id] = texts_by_q

        student_list = [(sub.student_id, sub.student.name) for sub in subs]
        for i in range(len(student_list)):
            for j in range(i + 1, len(student_list)):
                id_a, name_a = student_list[i]
                id_b, name_b = student_list[j]
                ta, tb = essay_texts.get(id_a, {}), essay_texts.get(id_b, {})
                suspicious_qs = []
                for eq in essay_questions:
                    text_a = ta.get(eq.id, "")
                    text_b = tb.get(eq.id, "")
                    if len(text_a) < 30 or len(text_b) < 30:
                        continue
                    ratio = difflib.SequenceMatcher(None, text_a, text_b).ratio()
                    if ratio >= 0.65:
                        suspicious_qs.append((eq.order + 1, round(ratio, 2)))
                if suspicious_qs:
                    avg_sim = round(sum(r for _, r in suspicious_qs) / len(suspicious_qs), 2)
                    labels = [f"Q{order} ({int(r*100)}%)" for order, r in suspicious_qs[:5]]
                    plagiarism_alerts.append({
                        "student_a": {"id": id_a, "name": name_a},
                        "student_b": {"id": id_b, "name": name_b},
                        "type": "essay",
                        "similarity": avg_sim,
                        "evidence": f"Similaridade alta nas dissertativas: {', '.join(labels)}",
                    })

    plagiarism_alerts.sort(key=lambda x: x["similarity"], reverse=True)

    return {
        "exam_id": exam_id,
        "exam_title": exam.title,
        "total_points": total_points,
        "total_enrolled": total_enrolled,
        "total_submitted": total_submitted,
        "corrected_count": len(subs),
        "average_score": avg_score,
        "median_score": median_score,
        "highest_score": highest,
        "lowest_score": lowest,
        "std_dev": std_dev,
        "pass_rate": pass_rate,
        "score_distribution": score_distribution,
        "question_stats": question_stats,
        "plagiarism_alerts": plagiarism_alerts,
    }


@router.get("/ai/status")
def ai_status(_: User = Depends(get_current_user)):
    available = groq_is_available()
    return {
        "available": available,
        "model": settings.GROQ_MODEL,
        "message": f"IA Online ({settings.GROQ_MODEL}) ativa" if available else "GROQ_API_KEY não configurada",
    }


@router.get("/ollama/status")
def ollama_status(_: User = Depends(get_current_user)):
    available = groq_is_available()
    return {
        "available": available,
        "model": settings.GROQ_MODEL,
        "message": f"IA Online ({settings.GROQ_MODEL}) ativa" if available else "GROQ_API_KEY não configurada",
    }
